from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Optional, Union
from zipfile import ZipFile


HOST_ROOT = Path(__file__).resolve().parent
PROCESSED_DIR = HOST_ROOT / "prompt" / "processed"
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-flash-latest")


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="TRPGシナリオを日本語GMプロンプトへ変換します。")
    parser.add_argument("source", help="入力ファイル。txt/docx/pdf/doc をサポートします。")
    parser.add_argument("--out-dir", default=str(PROCESSED_DIR), help="出力先ディレクトリ。")
    parser.add_argument("--no-gemini", action="store_true", help="Geminiを呼ばず、抽出テキストだけで出力します。")
    args = parser.parse_args(argv)

    source = Path(args.source).resolve()
    out_dir = Path(args.out_dir).resolve()
    text = extract_text(source)
    if not text.strip():
        raise SystemExit(f"テキストを抽出できませんでした: {source}")

    if args.no_gemini:
        scenario = fallback_structured_scenario(source, text)
    else:
        scenario = convert_with_gemini(source, text)

    out_dir.mkdir(parents=True, exist_ok=True)
    base = safe_stem(source)
    json_path = out_dir / f"{base}.json"
    md_path = out_dir / f"{base}.md"
    json_path.write_text(json.dumps(scenario, ensure_ascii=False, indent=2), encoding="utf-8")
    md_path.write_text(render_markdown(scenario), encoding="utf-8")
    print(f"JSON: {json_path}")
    print(f"MD: {md_path}")
    return 0


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return read_text_file(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".doc":
        converted = convert_doc_to_docx(path)
        try:
            return extract_docx_text(converted)
        finally:
            converted.unlink(missing_ok=True)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def read_text_file(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp932", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        table_cells: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        table_cells.append(cell.text.strip())
        return "\n".join(paragraphs + table_cells)
    except ImportError:
        return extract_docx_text_without_dependency(path)


def extract_docx_text_without_dependency(path: Path) -> str:
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)).strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def extract_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PDF抽出には pypdf が必要です。") from exc
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def convert_doc_to_docx(path: Path) -> Path:
    if sys.platform != "win32":
        raise RuntimeError(".doc変換はWindows上のMicrosoft Word COMを使用します。")
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise RuntimeError(".doc変換には pywin32 とMicrosoft Wordが必要です。") from exc

    output = Path(tempfile.gettempdir()) / f"{path.stem}.converted.docx"
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    document = None
    try:
        document = word.Documents.Open(str(path))
        document.SaveAs(str(output), FileFormat=16)
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()
    return output


def convert_with_gemini(source: Path, text: str) -> dict[str, Any]:
    api_key = load_gemini_api_key()
    if not api_key:
        raise RuntimeError("Gemini API keyが見つかりません。GEMINI_API_KEYまたはhost/gemini_api_key.*を設定してください。")
    try:
        import google.generativeai as genai  # type: ignore
    except ImportError as exc:
        raise RuntimeError("Gemini変換には google-generativeai が必要です。") from exc

    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(GEMINI_MODEL)
    prompt = build_conversion_prompt(source.name, text)
    response = model.generate_content(prompt)
    raw = getattr(response, "text", "") or ""
    parsed = parse_json_response(raw)
    if not parsed:
        raise RuntimeError("Gemini応答からJSONを解析できませんでした。")
    return normalize_scenario(source, parsed)


def build_conversion_prompt(filename: str, text: str) -> str:
    clipped = text[:60000]
    return (
        "あなたはTRPGシナリオ編集者です。入力された中国語または英語のシナリオを、"
        "日本語でローカルLLM用GMプロンプト素材に整理してください。\n"
        "必ずJSONだけを返してください。Markdownフェンスや説明文は禁止です。\n"
        "JSON schema:\n"
        "{\n"
        '  "title": "日本語タイトル",\n'
        '  "summary": "短い概要",\n'
        '  "gm_prompt": "GMが参照する日本語プロンプト本文",\n'
        '  "npcs": [{"name": "...", "description": "..."}],\n'
        '  "locations": [{"name": "...", "description": "..."}],\n'
        '  "scenes": [{"title": "...", "description": "...", "goals": []}],\n'
        '  "rules": ["進行上の注意"],\n'
        '  "initial_state_hints": {"current_scene": "...", "inventory": []}\n'
        "}\n"
        f"filename: {filename}\n"
        "source text:\n"
        f"{clipped}"
    )


def normalize_scenario(source: Path, value: dict[str, Any]) -> dict[str, Any]:
    return {
        "source_file": str(source),
        "title": str(value.get("title") or source.stem),
        "summary": str(value.get("summary") or ""),
        "gm_prompt": str(value.get("gm_prompt") or value.get("summary") or ""),
        "npcs": _list_of_dicts(value.get("npcs")),
        "locations": _list_of_dicts(value.get("locations")),
        "scenes": _list_of_dicts(value.get("scenes")),
        "rules": _list_of_text(value.get("rules")),
        "initial_state_hints": value.get("initial_state_hints") if isinstance(value.get("initial_state_hints"), dict) else {},
    }


def fallback_structured_scenario(source: Path, text: str) -> dict[str, Any]:
    return {
        "source_file": str(source),
        "title": source.stem,
        "summary": text[:500],
        "gm_prompt": text,
        "npcs": [],
        "locations": [],
        "scenes": [{"title": "開始", "description": text[:1200], "goals": []}],
        "rules": ["プレイヤーの行動を代行せず、自然な日本語で進行する。"],
        "initial_state_hints": {"current_scene": "開始", "inventory": []},
    }


def render_markdown(scenario: dict[str, Any]) -> str:
    lines = [
        f"# {scenario['title']}",
        "",
        "## 概要",
        scenario.get("summary", ""),
        "",
        "## GMプロンプト",
        scenario.get("gm_prompt", ""),
        "",
    ]
    for key, title in (("npcs", "NPC"), ("locations", "場所"), ("scenes", "シーン")):
        lines.extend([f"## {title}", ""])
        values = scenario.get(key) or []
        if not values:
            lines.append("- なし")
        for item in values:
            name = item.get("name") or item.get("title") or "未設定"
            description = item.get("description") or ""
            lines.append(f"- **{name}**: {description}")
        lines.append("")
    lines.extend(["## ルール", ""])
    for rule in scenario.get("rules") or []:
        lines.append(f"- {rule}")
    return "\n".join(lines).strip() + "\n"


def parse_json_response(text: str) -> Optional[dict[str, Any]]:
    cleaned = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.IGNORECASE | re.MULTILINE).strip()
    candidates = [cleaned]
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start != -1 and end != -1 and start < end:
        candidates.insert(0, cleaned[start : end + 1])
    for candidate in candidates:
        try:
            parsed = json.loads(candidate)
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            return parsed
    return None


def load_gemini_api_key() -> str:
    env_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if env_key:
        return env_key
    for path in (
        HOST_ROOT / "gemini_api_key.txt",
        HOST_ROOT / "gemini_api_key.py",
        HOST_ROOT / "gemini_apikey.txt",
        HOST_ROOT / "auto_tagger.py",
    ):
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        if path.suffix == ".py":
            match = re.search(r"GEMINI_API_KEY\s*=\s*(?:os\.environ\.get\([^,]+,\s*)?[\"']([^\"']+)[\"']", text)
            if match:
                return match.group(1).strip()
        elif text.strip():
            return text.strip()
    return ""


def safe_stem(path: Path) -> str:
    return re.sub(r"[^\w\-.一-龥ぁ-んァ-ンー]+", "_", path.stem, flags=re.UNICODE).strip("_") or "scenario"


def _list_of_dicts(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _list_of_text(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if not isinstance(value, list):
        return []
    return [str(item) for item in value if str(item).strip()]


if __name__ == "__main__":
    raise SystemExit(main())
